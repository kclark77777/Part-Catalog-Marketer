import pandas as pd
from docx import Document
import streamlit as st

# Load Excel file
def load_data(file_path):
    return pd.ExcelFile(file_path)

# Read relevant sheets
def filter_data(excel_data, selected_aircraft):
    parts_df = excel_data.parse('Parts')  # Sheet containing parts
    mro_df = excel_data.parse('MRO')      # Sheet containing MRO capabilities
    
    # Filter based on selected aircraft models
    filtered_parts = parts_df[parts_df['Aircraft Model'].isin(selected_aircraft)]
    filtered_mro = mro_df[mro_df['Aircraft Model'].isin(selected_aircraft)]
    
    return filtered_parts, filtered_mro

# Generate a Word document
def generate_document(selected_aircraft, parts, mro):
    doc = Document()
    doc.add_heading('Sales Collateral', 0)
    doc.add_paragraph(f'Selected Aircraft: {", ".join(selected_aircraft)}')
    
    doc.add_heading('Available Parts', level=1)
    for _, row in parts.iterrows():
        doc.add_paragraph(f"- {row['Part Number']}: {row['Description']}")
    
    doc.add_heading('MRO Capabilities', level=1)
    for _, row in mro.iterrows():
        doc.add_paragraph(f"- {row['Capability']} (Location: {row['Facility']})")
    
    file_path = 'Sales_Collateral.docx'
    doc.save(file_path)
    return file_path

# Streamlit Web App
def main():
    st.title("Aircraft Sales Collateral Generator")
    
    uploaded_file = st.file_uploader("Upload Excel File", type=["xlsx"])
    
    if uploaded_file:
        excel_data = load_data(uploaded_file)
        aircraft_models = list(excel_data.parse('Parts')['Aircraft Model'].unique())
        
        selected_aircraft = st.multiselect("Select Aircraft Models:", aircraft_models)
        
        if st.button("Generate Document"):
            if selected_aircraft:
                parts, mro = filter_data(excel_data, selected_aircraft)
                doc_path = generate_document(selected_aircraft, parts, mro)
                st.success("Document generated successfully!")
                with open(doc_path, "rb") as file:
                    st.download_button("Download Sales Collateral", file, file_name=doc_path)
            else:
                st.warning("Please select at least one aircraft model.")

if __name__ == "__main__":
    main()
